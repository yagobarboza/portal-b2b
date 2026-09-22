"""Ingestão de estoque — AGENTE (push), ARQUIVO (CSV/Excel) e WEBHOOK.

Três origens convergem para o MESMO motor de aplicação (`apply_stock_sync`):

1. AGENTE (Bloco I1, tipo `agent`) — push via chave de API:
   POST /integrations/agent/stock
2. ARQUIVO (Bloco B2, tipo `file`) — importação de CSV/Excel:
   POST /integrations/{integration_id}/stock/import
3. WEBHOOK (Bloco B3, tipo `webhook`) — evento `stock.sync` do ERP:
   POST /webhooks/{integration_id}  (assinatura HMAC-SHA256)

Garantias (comuns às três origens):
- Isolamento por tenant: o tenant vem da CHAVE (agente), da SESSÃO (arquivo)
  ou da INTEGRAÇÃO resolvida pela URL assinada (webhook).
- Idempotência: inbox transacional no PostgreSQL antes de chamar este motor.
- Estoque SEMPRE inteiro e ≥ 0 (truncado — nunca "10.000").
- Nunca cria produto: SKU inexistente vira erro tratado (catálogo curado).
- Cada execução fica registrada em SyncExecution (processed/errors/mensagem).
- N+1 evitado: SKUs são resolvidos e atualizados em chunks de 500.
"""
import csv
import io
from datetime import datetime, timezone

from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import bindparam, or_, select, update
from sqlalchemy.ext.asyncio import AsyncSession

from app.integrations.adapters import MappingStockAdapter
from app.integrations.contracts import StockUpdate, normalize_sku
from app.integrations.interfaces import StockAdapter
from app.models import Product, SyncExecution
from app.models.enums import SyncStatus
from app.repositories.catalog import invalidate_product_cache
from app.schemas.integration import MAX_IMPORT_ROWS

# Entidade registrada em SyncExecution para ingestão de estoque.
ENTITY = "stock"
# Teto de mensagens devolvidas ao importador (anti-payload gigante).
MAX_DETAILS = 50
STOCK_CHUNK_SIZE = 500

# Cabeçalhos aceitos no arquivo (case-insensitive) — Bloco B2.
SKU_HEADERS = ("sku", "codigo", "código", "codigo_sku")
STOCK_HEADERS = ("stock", "estoque", "quantidade", "qtd", "saldo")
EXTERNAL_ID_HEADERS = ("external_id", "id_externo")
OCCURRED_AT_HEADERS = ("occurred_at", "data_hora", "atualizado_em")
SOURCE_VERSION_HEADERS = ("source_version", "versao", "versão")

def _message(processed: int, unchanged: int, stale: int, errors: int) -> str:
    return (
        f"{processed} atualizado(s), {unchanged} inalterado(s), "
        f"{stale} obsoleto(s), {errors} erro(s)."
    )

def _first_error(exc: ValidationError) -> str:
    """Extrai a mensagem legível do erro de validação de um item."""
    try:
        msg = str(exc.errors()[0].get("msg", "Valor inválido."))
    except Exception:  # noqa: BLE001
        return "Valor inválido."
    return msg.replace("Value error, ", "")

# ==================== BLOCO B2 — Leitura do arquivo ====================
def _norm_key(value) -> str:
    """Normaliza o nome de uma coluna (case-insensitive, sem espaços)."""
    return str(value).strip().lower() if value is not None else ""

def _sniff_delimiter(sample: str) -> str:
    """Detecta ',' ou ';' no CSV (exportações pt-BR costumam usar ';')."""
    first = sample.splitlines()[0] if sample and sample.strip() else ""
    return ";" if first.count(";") > first.count(",") else ","

def iter_stock_rows(filename: str, content: bytes):
    """Gera linhas (dict) a partir de CSV ou Excel (.xlsx/.xlsm).

    - .xlsx/.xlsm: primeira planilha, primeira linha = cabeçalho.
    - Demais: interpretado como CSV (UTF-8 com BOM), com detecção de ';' ou ','.
    A chave de cada dict é o nome da coluna em minúsculas.
    """
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        headers = [_norm_key(h) for h in next(rows, ())]
        for values in rows:
            yield {h: ("" if v is None else v) for h, v in zip(headers, values) if h}
        wb.close()
        return

    text = content.decode("utf-8-sig")
    delimiter = _sniff_delimiter(text)
    for row in csv.DictReader(io.StringIO(text), delimiter=delimiter):
        yield {
            _norm_key(k): ("" if v is None else v)
            for k, v in row.items()
            if k
        }

def _pick(row: dict, headers: tuple[str, ...]) -> str:
    """Retorna o primeiro valor não-vazio entre os aliases de cabeçalho."""
    for header in headers:
        value = row.get(header)
        if value is not None and str(value).strip() != "":
            return str(value).strip()
    return ""

def parse_stock_rows(
    rows, max_rows: int = MAX_IMPORT_ROWS
) -> tuple[list[StockUpdate], list[dict]]:
    """Converte as linhas do arquivo em StockUpdate, separando erros por linha.

    NUNCA levanta por linha inválida: cada problema vira um item em `errors`
    (com o número da linha do arquivo) e o restante continua sendo importado.
    A primeira linha (cabeçalho) corresponde à linha 1, então os dados
    começam na linha 2.
    """
    items: list[StockUpdate] = []
    errors: list[dict] = []
    data_rows = 0
    adapter = MappingStockAdapter()

    for line, row in enumerate(rows, start=2):
        data_rows += 1
        if data_rows > max_rows:
            errors.append(
                {"row": line, "error": f"Arquivo excede o limite de {max_rows} linhas."}
            )
            break

        sku = _pick(row, SKU_HEADERS)
        if not sku:
            errors.append({"row": line, "error": "Linha sem SKU."})
            continue

        raw_stock = _pick(row, STOCK_HEADERS)
        if raw_stock == "":
            errors.append({"row": line, "sku": sku, "error": "Linha sem valor de estoque."})
            continue

        try:
            item = adapter.adapt_stock(
                {
                    "sku": sku,
                    "stock": raw_stock,
                    "external_id": _pick(row, EXTERNAL_ID_HEADERS) or None,
                    "occurred_at": _pick(row, OCCURRED_AT_HEADERS) or None,
                    "source_version": _pick(row, SOURCE_VERSION_HEADERS) or None,
                }
            )
        except ValidationError as exc:
            errors.append({"row": line, "sku": sku, "error": _first_error(exc)})
            continue

        items.append(item)

    if data_rows == 0:
        errors.append(
            {"row": 1, "error": "Nenhuma linha de dados encontrada (apenas o cabeçalho)."}
        )
    return items, errors

async def apply_stock_import(
    db: AsyncSession,
    *,
    integration,
    filename: str,
    content: bytes,
    sync_execution: SyncExecution | None = None,
) -> dict:
    """Importa o estoque de um arquivo (CSV/Excel) e registra a execução.

    Reusa integralmente `apply_stock_sync`: os erros de PARSING do arquivo
    entram como `extra_errors` e são somados aos erros de aplicação.
    """
    items, errors = parse_stock_rows(iter_stock_rows(filename, content))
    return await apply_stock_sync(
        db,
        integration=integration,
        items=items,
        batch_id=None,
        extra_errors=errors,
        sync_execution=sync_execution,
    )

# ==================== BLOCO B3 — Registros do webhook ====================
def parse_stock_records(
    records: list[dict],
    max_records: int = MAX_IMPORT_ROWS,
    adapter: StockAdapter | None = None,
) -> tuple[list[StockUpdate], list[dict]]:
    """Converte registros externos em StockUpdate via adapter canônico.

    Espelha `parse_stock_rows`, mas para o payload do webhook: o campo
    `index` identifica a posição do registro no array (1-based). Um registro
    inválido NUNCA derruba o lote — vira item em `errors`, e o restante
    continua sendo aplicado.
    """
    items: list[StockUpdate] = []
    errors: list[dict] = []
    stock_adapter = adapter or MappingStockAdapter()

    for index, rec in enumerate(records, start=1):
        if index > max_records:
            errors.append(
                {"index": index, "error": f"Evento excede o limite de {max_records} registros."}
            )
            break
        if not isinstance(rec, dict):
            errors.append({"index": index, "error": "Registro inválido (não é objeto)."})
            continue

        source_sku = str(rec.get("sku") or "").strip()
        try:
            item = stock_adapter.adapt_stock(rec)
        except ValidationError as exc:
            detail = {"index": index, "error": _first_error(exc)}
            if source_sku:
                detail["sku"] = source_sku
            errors.append(detail)
            continue

        items.append(item)

    return items, errors

# ==================== Motor comum de aplicação ====================
async def apply_stock_sync(
    db: AsyncSession,
    *,
    integration,
    items,
    batch_id: str | None = None,
    extra_errors: list[dict] | None = None,
    sync_execution: SyncExecution | None = None,
) -> dict:
    """Aplica um lote de estoque e registra a execução.

    `extra_errors` (Blocos B2/B3): erros de PARSING (linha/registro inválido,
    SKU ausente, estoque inválido). Já vêm com o identificador da origem e são
    somados aos erros de aplicação, sem interromper o resto do lote.

    Retorna um dict pronto para `StockSyncResult`:
    {sync_id, status, processed, unchanged, errors, message, details}
    """
    tenant_id = integration.tenant_id
    now = datetime.now(timezone.utc)

    # A idempotência é garantida antes daqui pela inbox transacional.
    sync = sync_execution or SyncExecution(
        tenant_id=tenant_id,
        integration_id=integration.id,
        entity=ENTITY,
    )
    if sync_execution is None:
        db.add(sync)
    sync.status = SyncStatus.RUNNING
    sync.started_at = sync.started_at or now
    sync.last_attempt_at = now
    await db.flush()

    # Deduplica pelo evento temporal mais novo; no empate, a última ocorrência vence.
    wanted: dict[str, StockUpdate] = {}
    for item in items:
        key = normalize_sku(item.sku)
        previous = wanted.get(key)
        if key and (
            previous is None
            or item.occurred_at is None
            or previous.occurred_at is None
            or item.occurred_at >= previous.occurred_at
        ):
            wanted[key] = item

    found: dict[str, Product] = {}
    keys = list(wanted)
    for offset in range(0, len(keys), STOCK_CHUNK_SIZE):
        chunk = keys[offset : offset + STOCK_CHUNK_SIZE]
        result = await db.execute(
            select(Product).where(
                Product.tenant_id == tenant_id,
                Product.is_deleted == False,  # noqa: E712 — ignora soft delete
                Product.normalized_sku.in_(chunk),
            )
        )
        found.update({p.normalized_sku: p for p in result.scalars().all()})

    processed = 0
    unchanged = 0
    stale = 0
    errors = len(extra_errors or [])
    details: list[dict] = list(extra_errors or [])[:MAX_DETAILS]
    updates: list[dict] = []

    for key, item in wanted.items():
        product = found.get(key)
        if product is None:
            errors += 1
            if len(details) < MAX_DETAILS:
                details.append(
                    {"sku": item.sku, "error": "SKU não encontrado no catálogo."}
                )
            continue
        incoming_at = item.occurred_at or now
        if incoming_at.tzinfo is None:
            incoming_at = incoming_at.replace(tzinfo=timezone.utc)
        stored_at = product.stock_updated_at
        if stored_at is not None and stored_at.tzinfo is None:
            stored_at = stored_at.replace(tzinfo=timezone.utc)
        if item.source_version and item.source_version == product.stock_source_version:
            if product.stock == item.stock:
                unchanged += 1
            else:
                errors += 1
                stale += 1
                if len(details) < MAX_DETAILS:
                    details.append(
                        {"sku": item.sku, "error": "Versão da fonte repetida com saldo diferente."}
                    )
            continue
        if item.occurred_at is not None and stored_at is not None and incoming_at <= stored_at:
            stale += 1
            continue
        if product.stock == item.stock:
            unchanged += 1
            continue
        updates.append(
            {
                "_product_id": product.id,
                "_stock": item.stock,
                "_stock_updated_at": incoming_at,
                "_stock_updated_at_guard": incoming_at,
                "_stock_source_version": item.source_version,
            }
        )
        processed += 1

    statement = (
        update(Product.__table__)
        .where(
            Product.__table__.c.id == bindparam("_product_id"),
            or_(
                Product.__table__.c.stock_updated_at.is_(None),
                Product.__table__.c.stock_updated_at
                < bindparam("_stock_updated_at_guard"),
            ),
        )
        .values(
            stock=bindparam("_stock"),
            stock_updated_at=bindparam("_stock_updated_at"),
            stock_source_version=bindparam("_stock_source_version"),
        )
    )
    for offset in range(0, len(updates), STOCK_CHUNK_SIZE):
        await db.execute(statement, updates[offset : offset + STOCK_CHUNK_SIZE])

    # 5) Fecha a execução com o resumo (seção 33).
    if errors == 0:
        final_status = SyncStatus.SUCCESS
        label = "ok"
    elif processed or unchanged:
        final_status = SyncStatus.PARTIAL
        label = "partial"
    else:
        final_status = SyncStatus.FAILED
        label = "failed"

    sync.status = final_status
    sync.processed = processed
    sync.errors = errors
    sync.finished_at = datetime.now(timezone.utc)
    sync.terminal_at = sync.finished_at
    sync.next_retry_at = None
    sync.message = _message(processed, unchanged, stale, errors)
    result = {
        "sync_id": sync.id,
        "status": label,
        "processed": processed,
        "unchanged": unchanged,
        "stale": stale,
        "errors": errors,
        "message": sync.message,
        "details": details,
    }
    from app.services.integration_observability import finish_run_from_result

    finish_run_from_result(
        sync,
        result,
        items_received=len(items) + len(extra_errors or []),
    )
    await db.flush()
    if processed:
        await invalidate_product_cache(tenant_id)

    result["message"] = sync.message
    result["details"] = sync.item_errors or []
    return result
