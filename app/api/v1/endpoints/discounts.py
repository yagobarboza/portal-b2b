"""Endpoints de Desconto por Quantidade (Desconto Progressivo).

- GET/POST /catalog/discounts          -> listar/criar regras do tenant
- PATCH/DELETE /catalog/discounts/{id} -> editar/remover regra
- POST /catalog/discounts/import       -> importação em massa (CSV/Excel)
- Isolamento por tenant + RBAC (catalogs:manage)
"""
import csv
import io
import re
from decimal import Decimal
from uuid import UUID

from fastapi import APIRouter, Depends, File, Query, UploadFile
from openpyxl import load_workbook
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import require_permission
from app.core.exceptions import NotFoundError, ValidationError, ValidationFailedError
from app.core.permissions import CATALOG_MANAGE
from app.database.session import get_db
from app.models import Customer, Product, User
from app.repositories.catalog import ProductRepository
from app.repositories.discount import QuantityDiscountRepository
from app.schemas.discount import (
    QuantityDiscountCreate,
    QuantityDiscountImportResult,
    QuantityDiscountPage,
    QuantityDiscountRead,
    QuantityDiscountUpdate,
)
from app.services.audit import record_audit

router = APIRouter(prefix="/catalog/discounts", tags=["Descontos"])

# Limite anti-DoS para importação (mesmo padrão de /customers/import)
IMPORT_MAX_BYTES = 2 * 1024 * 1024  # 2 MB

_VALID_TYPES = {"percent", "fixed"}

def _normalize_digits(value: str) -> str:
    """Remove tudo que não for dígito (CPF/CNPJ)."""
    return re.sub(r"\D", "", value or "")

def _iter_discount_rows(filename: str, content: bytes):
    """Gera linhas (dict) a partir de CSV ou Excel (.xlsx/.xlsm).

    - .xlsx/.xlsm: lê a primeira planilha, primeira linha = cabeçalho.
    - Qualquer outro: tenta como CSV (UTF-8 com BOM).
    """
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        headers = [str(h).strip() if h is not None else "" for h in next(rows, ())]
        for values in rows:
            yield dict(zip(headers, ("" if v is None else v for v in values)))
        wb.close()
        return
    text = content.decode("utf-8-sig")
    yield from csv.DictReader(io.StringIO(text))

def _parse_decimal(raw) -> Decimal | None:
    """Converte valor aceitando vírgula decimal (ex.: '5,5' -> 5.5)."""
    try:
        return Decimal(str(raw or "").strip().replace(",", "."))
    except Exception:  # noqa: BLE001
        return None

def _parse_positive_int(raw) -> int | None:
    """Converte para inteiro > 0 (aceita '10' ou 10.0)."""
    try:
        value = int(Decimal(str(raw or "").strip().replace(",", ".")))
    except Exception:  # noqa: BLE001
        return None
    return value if value > 0 else None

def _validate_rule(discount_type: str, discount_value: Decimal) -> None:
    """Valida tipo/valor: percentual não pode passar de 100%."""
    if discount_type not in _VALID_TYPES:
        raise ValidationError("Tipo de desconto deve ser 'percent' ou 'fixed'.")
    if discount_value <= 0:
        raise ValidationError("Valor do desconto deve ser maior que zero.")
    if discount_type == "percent" and discount_value > Decimal("100"):
        raise ValidationError("Desconto percentual deve ser até 100%.")

async def _belongs_to_tenant(
    db: AsyncSession, model, obj_id: UUID, tenant_id
) -> bool:
    """Confirma que um registro (Customer/Product) pertence ao tenant.

    Bloqueia IDOR/BOLA: impede regra de desconto apontando para produto
    ou cliente de OUTRA empresa. Sem tenant → nunca autoriza.
    """
    if tenant_id is None:
        return False
    result = await db.execute(
        select(model.id).where(model.id == obj_id, model.tenant_id == tenant_id)
    )
    return result.scalars().first() is not None

async def _enrich(
    db: AsyncSession, discounts: list
) -> list[QuantityDiscountRead]:
    """Enriquece regras com nome/sku do produto e nome do cliente.

    Usa 2 queries em lote (IN ...) — nada de N+1 (requisito de volumetria).
    """
    product_ids = {d.product_id for d in discounts}
    customer_ids = {d.customer_id for d in discounts if d.customer_id}

    product_map: dict[UUID, tuple[str, str]] = {}
    if product_ids:
        rows = (
            await db.execute(
                select(Product.id, Product.name, Product.sku).where(
                    Product.id.in_(product_ids)
                )
            )
        ).all()
        product_map = {r[0]: (r[1], r[2]) for r in rows}

    customer_map: dict[UUID, str] = {}
    if customer_ids:
        rows = (
            await db.execute(
                select(Customer.id, Customer.name).where(
                    Customer.id.in_(customer_ids)
                )
            )
        ).all()
        customer_map = {r[0]: r[1] for r in rows}

    result: list[QuantityDiscountRead] = []
    for d in discounts:
        pname, psku = product_map.get(d.product_id, (None, None))
        dtype = (
            d.discount_type.value
            if hasattr(d.discount_type, "value")
            else d.discount_type
        )
        result.append(
            QuantityDiscountRead(
                id=d.id,
                product_id=d.product_id,
                product_name=pname,
                product_sku=psku,
                customer_id=d.customer_id,
                customer_name=(
                    customer_map.get(d.customer_id) if d.customer_id else None
                ),
                min_quantity=d.min_quantity,
                discount_type=dtype,
                discount_value=d.discount_value,
                is_active=d.is_active,
                created_at=d.created_at,
            )
        )
    return result

# ---------- Listagem ----------
@router.get("", response_model=QuantityDiscountPage)
async def list_discounts(
    product_id: UUID | None = None,
    customer_id: UUID | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),  # limite anti-DoS (seção 53)
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_permission(CATALOG_MANAGE)),
) -> QuantityDiscountPage:
    """Lista regras de desconto do tenant (filtros opcionais produto/cliente)."""
    repo = QuantityDiscountRepository(db)
    items, total = await repo.list_rules(product_id, customer_id, page, page_size)
    pages = (total + page_size - 1) // page_size
    enriched = await _enrich(db, items)
    return QuantityDiscountPage(
        items=enriched, total=total, page=page, page_size=page_size, pages=pages
    )

# ---------- Criação (manual) ----------
@router.post("", response_model=QuantityDiscountRead, status_code=201)
async def create_discount(
    body: QuantityDiscountCreate,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_permission(CATALOG_MANAGE)),
) -> QuantityDiscountRead:
    """Cria uma regra de desconto por quantidade (cadastro manual)."""
    tenant = user.tenant_id
    # Segurança: produto e cliente precisam pertencer ao MESMO tenant
    if not await _belongs_to_tenant(db, Product, body.product_id, tenant):
        raise NotFoundError("Produto não encontrado.")
    if body.customer_id and not await _belongs_to_tenant(
        db, Customer, body.customer_id, tenant
    ):
        raise NotFoundError("Cliente não encontrado.")

    _validate_rule(body.discount_type, body.discount_value)

    repo = QuantityDiscountRepository(db)
    existing = await repo.find_by_key(
        body.product_id, body.customer_id, body.min_quantity
    )
    if existing:
        raise ValidationFailedError(
            "Já existe uma regra para este produto/cliente nesta faixa de quantidade."
        )

    discount = await repo.create(body.model_dump())
    await record_audit(
        db, action="create", entity="quantity_discount",
        entity_id=discount.id, user_id=user.id, tenant_id=user.tenant_id,
    )
    await db.commit()
    enriched = await _enrich(db, [discount])
    return enriched[0]

# ---------- Edição ----------
@router.patch("/{discount_id}", response_model=QuantityDiscountRead)
async def update_discount(
    discount_id: UUID,
    body: QuantityDiscountUpdate,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_permission(CATALOG_MANAGE)),
) -> QuantityDiscountRead:
    """Edita uma regra de desconto (mesmo tenant obrigatório)."""
    repo = QuantityDiscountRepository(db)
    discount = await repo.get(discount_id)
    if not discount:
        raise NotFoundError("Regra de desconto não encontrada.")

    data = body.model_dump(exclude_unset=True)
    tenant = user.tenant_id

    if "product_id" in data and not await _belongs_to_tenant(
        db, Product, data["product_id"], tenant
    ):
        raise NotFoundError("Produto não encontrado.")
    if data.get("customer_id") and not await _belongs_to_tenant(
        db, Customer, data["customer_id"], tenant
    ):
        raise NotFoundError("Cliente não encontrado.")

    # Valida o par final (tipo/valor) após mesclar com o existente
    final_type = data.get("discount_type") or (
        discount.discount_type.value
        if hasattr(discount.discount_type, "value")
        else discount.discount_type
    )
    final_value = data.get("discount_value") or discount.discount_value
    _validate_rule(final_type, final_value)

    await repo.update(discount, data)
    await record_audit(
        db, action="update", entity="quantity_discount",
        entity_id=discount.id, user_id=user.id, tenant_id=user.tenant_id,
    )
    await db.commit()
    enriched = await _enrich(db, [discount])
    return enriched[0]

# ---------- Remoção ----------
@router.delete("/{discount_id}", status_code=204)
async def delete_discount(
    discount_id: UUID,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_permission(CATALOG_MANAGE)),
) -> None:
    """Remove uma regra de desconto (hard delete — a regra não tem soft delete)."""
    repo = QuantityDiscountRepository(db)
    discount = await repo.get(discount_id)
    if not discount:
        raise NotFoundError("Regra de desconto não encontrada.")
    await repo.delete(discount)
    await record_audit(
        db, action="delete", entity="quantity_discount",
        entity_id=discount_id, user_id=user.id, tenant_id=user.tenant_id,
    )
    await db.commit()

# ---------- Importação em massa (CSV/Excel) ----------
@router.post("/import", response_model=QuantityDiscountImportResult)
async def import_discounts(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_permission(CATALOG_MANAGE)),
) -> QuantityDiscountImportResult:
    """Importa regras de desconto em massa (CSV ou Excel).

    Colunas esperadas:
      - sku:             SKU do produto (obrigatório).
      - document:        CPF/CNPJ do cliente (OPCIONAL — vazio = todos).
      - min_quantity:    quantidade mínima da faixa (obrigatório, inteiro > 0).
      - discount_type:   'percent' | 'fixed' (obrigatório).
      - discount_value:  valor do desconto (aceita vírgula decimal).

    UPSERT: mesma chave (produto + cliente + faixa) → ATUALIZA a regra
    existente em vez de duplicar. Tudo em um único commit (volumetria).
    """
    content = await file.read()
    if len(content) > IMPORT_MAX_BYTES:
        raise ValidationError("Arquivo excede o limite de 2 MB.")

    tenant = user.tenant_id
    repo = QuantityDiscountRepository(db)
    product_repo = ProductRepository(db)

    created = updated = skipped = 0
    errors: list[dict] = []
    product_cache: dict[str, UUID | None] = {}
    customer_cache: dict[str, UUID | None] = {}

    async def _resolve_customer(document: str) -> UUID | None:
        """Resolve o cliente pelo documento (com ou sem formatação)."""
        if document in customer_cache:
            return customer_cache[document]
        digits = _normalize_digits(document)
        stmt = select(Customer.id).where(
            Customer.tenant_id == tenant,
            or_(
                Customer.document == document,
                Customer.document == digits,
                func.replace(
                    func.replace(
                        func.replace(Customer.document, ".", ""), "/", ""
                    ),
                    "-",
                    "",
                )
                == digits,
            ),
        )
        found = (await db.execute(stmt)).scalars().first()
        customer_cache[document] = found
        return found

    try:
        for idx, row in enumerate(
            _iter_discount_rows(file.filename, content), start=2
        ):
            sku = str(row.get("sku") or "").strip()
            document = str(row.get("document") or "").strip()
            min_quantity = _parse_positive_int(row.get("min_quantity"))
            discount_type = str(row.get("discount_type") or "").strip().lower()
            discount_value = _parse_decimal(row.get("discount_value"))

            # ---- Validações da linha (erro por linha, nunca aborta o lote) ----
            if not sku:
                skipped += 1
                errors.append({"row": idx, "error": "SKU é obrigatório."})
                continue
            if min_quantity is None:
                skipped += 1
                errors.append(
                    {"row": idx, "error": "min_quantity deve ser inteiro > 0."}
                )
                continue
            if discount_type not in _VALID_TYPES:
                skipped += 1
                errors.append(
                    {
                        "row": idx,
                        "error": "discount_type deve ser 'percent' ou 'fixed'.",
                    }
                )
                continue
            if discount_value is None or discount_value <= 0:
                skipped += 1
                errors.append({"row": idx, "error": "discount_value inválido."})
                continue
            if discount_type == "percent" and discount_value > Decimal("100"):
                skipped += 1
                errors.append(
                    {
                        "row": idx,
                        "error": "Desconto percentual deve ser até 100%.",
                    }
                )
                continue

            # ---- Resolve produto (cache em memória — volumetria) ----
            if sku not in product_cache:
                product = await product_repo.get_by_sku(sku)
                product_cache[sku] = product.id if product else None
            product_id = product_cache[sku]
            if not product_id:
                skipped += 1
                errors.append(
                    {"row": idx, "error": f"SKU '{sku}' não encontrado."}
                )
                continue

            # ---- Resolve cliente (vazio = regra global) ----
            customer_id: UUID | None = None
            if document:
                customer_id = await _resolve_customer(document)
                if not customer_id:
                    skipped += 1
                    errors.append(
                        {
                            "row": idx,
                            "error": f"Cliente '{document}' não encontrado.",
                        }
                    )
                    continue

            # ---- UPSERT pela chave (produto + cliente + faixa) ----
            existing = await repo.find_by_key(
                product_id, customer_id, min_quantity
            )
            payload = {
                "product_id": product_id,
                "customer_id": customer_id,
                "min_quantity": min_quantity,
                "discount_type": discount_type,
                "discount_value": discount_value,
                "is_active": True,
            }
            if existing:
                await repo.update(existing, payload)
                updated += 1
            else:
                await repo.create(payload)
                created += 1

        await db.commit()
    except Exception as exc:  # noqa: BLE001
        await db.rollback()
        raise ValidationFailedError(f"Falha ao processar o arquivo: {exc}")

    await record_audit(
        db, action="import", entity="quantity_discount",
        user_id=user.id, tenant_id=user.tenant_id,
    )
    await db.commit()
    return QuantityDiscountImportResult(
        created=created, updated=updated, skipped=skipped, errors=errors
    )