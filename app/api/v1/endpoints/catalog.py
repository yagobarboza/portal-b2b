"""Endpoints do catálogo (seções 16 e 17 do doc).

- Categorias: CRUD (Admin/Vendedor).
- Produtos: CRUD + busca/filtros/ordenação/paginação (Admin/Vendedor).
- Tabelas de preço e preço por cliente (seção 17).
- Cotação de preço: o backend recalcula o preço final (nunca confia no frontend).
- Importação em massa de preços especiais (CSV/Excel) — Bloco B3.
- Isolamento por tenant em todas as queries (seção 5).
- RBAC: require_permission (seção 13).
"""
import csv
import io
import re
from decimal import Decimal
from uuid import UUID

from fastapi import APIRouter, Depends, File, Query, UploadFile
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import get_current_user, require_permission
from app.core.exceptions import NotFoundError, ValidationFailedError
from app.core.permissions import CATALOG_MANAGE
from app.database.session import get_db
from app.models import Customer, CustomerPrice, Product, User
from app.repositories.catalog import (
    CategoryRepository,
    CustomerPriceRepository,
    PriceListRepository,
    ProductRepository,
)
from app.repositories.customer import CustomerRepository
from app.schemas.catalog import (
    CategoryCreate,
    CategoryRead,
    CategoryUpdate,
    CustomerPriceCreate,
    CustomerPriceDetailRead,
    CustomerPriceImportResult,
    CustomerPricePage,
    CustomerPriceRead,
    CustomerPriceUpdate,
    PriceListCreate,
    PriceListRead,
    PriceQuote,
    ProductCreate,
    ProductListParams,
    ProductPage,
    ProductRead,
    ProductUpdate,
)
from app.services.audit import record_audit
from app.services.pricing import calculate_product_price
from app.services.product_images import attach_product_images

router = APIRouter(prefix="/catalog", tags=["Catálogo"])

# Limite anti-DoS para importação (mesmo padrão do /customers/import)
IMPORT_MAX_BYTES = 2 * 1024 * 1024  # 2 MB

async def _belongs_to_tenant(
    db: AsyncSession, model, obj_id: UUID, tenant_id
) -> bool:
    """Confirma que um registro (Customer/Product) pertence ao tenant.

    Bloqueia IDOR/BOLA: impede criar preço especial apontando para
    cliente ou produto de OUTRA empresa. Sem tenant → nunca autoriza.
    """
    if tenant_id is None:
        return False
    result = await db.execute(
        select(model).where(
            model.id == obj_id,
            model.tenant_id == tenant_id,
        )
    )
    return result.scalars().first() is not None

def _resolve_quote_customer(user: User, requested: UUID | None) -> UUID | None:
    """Define qual customer_id entra no cálculo de preço (anti-vazamento).

    - Cliente do portal: SEMPRE o próprio id — nunca o de outro cliente.
    - Usuário da empresa (mesmo tenant): pode consultar clientes do tenant.
    - Super admin / sem tenant: nenhum preço especial.
    """
    if user.customer_id is not None:
        return user.customer_id
    if user.tenant_id is not None:
        return requested
    return None

def _normalize_digits(value: str) -> str:
    """Remove tudo que não for dígito (CPF/CNPJ)."""
    return re.sub(r"\D", "", value or "")

def _iter_price_rows(filename: str, content: bytes):
    """Gera linhas (dict) a partir de CSV ou Excel (.xlsx/.xlsm).

    - .xlsx/.xlsm: lê a primeira planilha, primeira linha = cabeçalho.
    - Qualquer outro: tenta como CSV (UTF-8 com BOM).
    """
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        headers = [str(h).strip() if h is not None else "" for h in next(rows, ())]
        for values in rows:
            yield dict(zip(headers, ("" if v is None else v for v in values)))
        wb.close()
        return
    text = content.decode("utf-8-sig")
    yield from csv.DictReader(io.StringIO(text))

# ---------- Categorias ----------
@router.post("/categories", response_model=CategoryRead, status_code=201)
async def create_category(
    body: CategoryCreate,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_permission(CATALOG_MANAGE)),
) -> CategoryRead:
    repo = CategoryRepository(db)
    category = await repo.create(body.model_dump())
    await record_audit(
        db, action="create", entity="category",
        entity_id=category.id, user_id=user.id, tenant_id=user.tenant_id,
    )
    await db.commit()
    return category

@router.get("/categories", response_model=list[CategoryRead])
async def list_categories(
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
) -> list[CategoryRead]:
    repo = CategoryRepository(db)
    return await repo.list_all()

@router.get("/categories/{category_id}", response_model=CategoryRead)
async def get_category(
    category_id: UUID,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
) -> CategoryRead:
    repo = CategoryRepository(db)
    category = await repo.get(category_id)
    if not category:
        raise NotFoundError("Categoria não encontrada.")
    return category

@router.patch("/categories/{category_id}", response_model=CategoryRead)
async def update_category(
    category_id: UUID,
    body: CategoryUpdate,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_permission(CATALOG_MANAGE)),
) -> CategoryRead:
    repo = CategoryRepository(db)
    category = await repo.get(category_id)
    if not category:
        raise NotFoundError("Categoria não encontrada.")
    data = {k: v for k, v in body.model_dump().items() if v is not None}
    for key, value in data.items():
        setattr(category, key, value)
    await record_audit(
        db, action="update", entity="category",
        entity_id=category.id, user_id=user.id, tenant_id=user.tenant_id,
    )
    await db.commit()
    return category

# ---------- Produtos ----------
@router.post("/products", response_model=ProductRead, status_code=201)
async def create_product(
    body: ProductCreate,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_permission(CATALOG_MANAGE)),
) -> ProductRead:
    repo = ProductRepository(db)
    product = await repo.create(body.model_dump())
    await record_audit(
        db, action="create", entity="product",
        entity_id=product.id, user_id=user.id, tenant_id=user.tenant_id,
    )
    await db.commit()
    return product

@router.get("/products", response_model=ProductPage)
async def list_products(
    params: ProductListParams = Depends(),
    customer_id: UUID | None = None,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
) -> ProductPage:
    """Lista produtos com busca, filtros, ordenação e paginação (seção 53).

    Quando o solicitante é um cliente (ou um usuário da empresa consulta um
    cliente do tenant), o backend anexa o preço calculado em 1 query em batch.
    """
    repo = ProductRepository(db)
    items, total = await repo.search(
        search=params.search,
        category_id=params.category_id,
        status=params.status,
        min_price=params.min_price,
        max_price=params.max_price,
        sort_by=params.sort_by,
        sort_dir=params.sort_dir,
        page=params.page,
        page_size=params.page_size,
    )
    pages = (total + params.page_size - 1) // params.page_size

    # Resolve a URL da imagem de cada produto em UMA query (R2 signed URL)
    image_urls = await attach_product_images(db, items)
    enriched = [
        ProductRead.model_validate(p).model_copy(
            update={"image_url": image_urls.get(p.id)}
        )
        for p in items
    ]

    # Preços especiais na vitrine (batch — sem N+1)
    effective_customer = _resolve_quote_customer(user, customer_id)
    if effective_customer is not None and items:
        price_map = await CustomerPriceRepository(db).get_many_for_customer(
            effective_customer, [p.id for p in items]
        )
        final_items: list[ProductRead] = []
        for pr, p in zip(enriched, items):
            cp = price_map.get(p.id)
            if cp is not None:
                final_items.append(
                    pr.model_copy(
                        update={
                            "customer_price": cp,
                            "final_price": cp,
                            "price_source": "customer",
                        }
                    )
                )
            else:
                final_items.append(
                    pr.model_copy(
                        update={"final_price": p.price, "price_source": "default"}
                    )
                )
    else:
        final_items = enriched

    return ProductPage(
        items=final_items,
        total=total,
        page=params.page,
        page_size=params.page_size,
        pages=pages,
    )

# ---------- Resolução por SKU (cadastro manual — Bloco B3) ----------
@router.get("/products/by-sku", response_model=ProductRead)
async def get_product_by_sku(
    sku: str = Query(..., min_length=1, max_length=80),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
) -> ProductRead:
    """Resolve um produto pelo SKU exato do tenant (cadastro manual por SKU)."""
    repo = ProductRepository(db)
    product = await repo.get_by_sku(sku)
    if not product:
        raise NotFoundError("Produto não encontrado para o SKU informado.")
    image_urls = await attach_product_images(db, [product])
    return ProductRead.model_validate(product).model_copy(
        update={"image_url": image_urls.get(product.id)}
    )

@router.get("/products/{product_id}", response_model=ProductRead)
async def get_product(
    product_id: UUID,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
) -> ProductRead:
    repo = ProductRepository(db)
    product = await repo.get(product_id)
    if not product:
        raise NotFoundError("Produto não encontrado.")

    image_urls = await attach_product_images(db, [product])
    return ProductRead.model_validate(product).model_copy(
        update={"image_url": image_urls.get(product.id)}
    )

@router.patch("/products/{product_id}", response_model=ProductRead)
async def update_product(
    product_id: UUID,
    body: ProductUpdate,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_permission(CATALOG_MANAGE)),
) -> ProductRead:
    repo = ProductRepository(db)
    product = await repo.get(product_id)
    if not product:
        raise NotFoundError("Produto não encontrado.")
    data = {k: v for k, v in body.model_dump().items() if v is not None}
    for key, value in data.items():
        setattr(product, key, value)
    await record_audit(
        db, action="update", entity="product",
        entity_id=product.id, user_id=user.id, tenant_id=user.tenant_id,
    )
    await db.commit()
    return product

# ---------- Tabelas de preço (seção 17) ----------
@router.post("/price-lists", response_model=PriceListRead, status_code=201)
async def create_price_list(
    body: PriceListCreate,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_permission(CATALOG_MANAGE)),
) -> PriceListRead:
    repo = PriceListRepository(db)
    price_list = await repo.create(body.model_dump())
    await record_audit(
        db, action="create", entity="price_list",
        entity_id=price_list.id, user_id=user.id, tenant_id=user.tenant_id,
    )
    await db.commit()
    return price_list

# ---------- Preço por cliente (seção 17) ----------
@router.post("/customer-prices", response_model=CustomerPriceRead, status_code=201)
async def create_customer_price(
    body: CustomerPriceCreate,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_permission(CATALOG_MANAGE)),
) -> CustomerPriceRead:
    repo = CustomerPriceRepository(db)
    # Segurança: cliente e produto precisam pertencer ao MESMO tenant
    tenant = user.tenant_id
    if not await _belongs_to_tenant(db, Customer, body.customer_id, tenant):
        raise NotFoundError("Cliente não encontrado.")
    if not await _belongs_to_tenant(db, Product, body.product_id, tenant):
        raise NotFoundError("Produto não encontrado.")
    try:
        cp = await repo.create(body.model_dump())
    except IntegrityError:
        await db.rollback()
        raise ValidationFailedError(
            "Já existe um preço especial para este cliente e produto."
        )
    await record_audit(
        db, action="create", entity="customer_price",
        entity_id=cp.id, user_id=user.id, tenant_id=user.tenant_id,
    )
    await db.commit()
    return cp

# ===== Gestão completa de preços especiais (Bloco A) =====
@router.get("/customer-prices", response_model=CustomerPricePage)
async def list_customer_prices(
    customer_id: UUID | None = None,
    product_id: UUID | None = None,
    search: str | None = Query(None, max_length=120),
    min_price: Decimal | None = Query(None, ge=0),
    max_price: Decimal | None = Query(None, ge=0),
    sort_by: str = Query("recent", pattern="^(recent|price_asc|price_desc)$"),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),  # limite anti-DoS (seção 53)
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_permission(CATALOG_MANAGE)),
) -> CustomerPricePage:
    """Lista preços especiais do tenant com busca, faixa de preço e ordenação.

    Apenas usuário com CATALOG_MANAGE; o repositório filtra por tenant.
    - search: nome do cliente, nome do produto ou SKU.
    - min_price / max_price: faixa do preço especial.
    - sort_by: recent | price_asc | price_desc.
    """
    repo = CustomerPriceRepository(db)
    items, total = await repo.list_all(
        customer_id=customer_id,
        product_id=product_id,
        search=search,
        min_price=min_price,
        max_price=max_price,
        sort_by=sort_by,
        page=page,
        page_size=page_size,
    )

    # Enriquecimento com nomes (batches — sem N+1)
    product_ids = {cp.product_id for cp in items}
    customer_ids = {cp.customer_id for cp in items}
    products: dict = {}
    customers: dict = {}
    if product_ids:
        for p in (await db.execute(
            select(Product).where(Product.id.in_(product_ids))
        )).scalars().all():
            products[p.id] = p
    if customer_ids:
        for c in (await db.execute(
            select(Customer).where(Customer.id.in_(customer_ids))
        )).scalars().all():
            customers[c.id] = c

    items_out = [
        CustomerPriceDetailRead(
            id=cp.id,
            customer_id=cp.customer_id,
            customer_name=customers[cp.customer_id].name if cp.customer_id in customers else None,
            product_id=cp.product_id,
            product_name=products[cp.product_id].name if cp.product_id in products else None,
            product_sku=products[cp.product_id].sku if cp.product_id in products else None,
            price=cp.price,
        )
        for cp in items
    ]
    pages = (total + page_size - 1) // page_size if total else 1
    return CustomerPricePage(
        items=items_out, total=total, page=page, page_size=page_size, pages=pages
    )

@router.patch("/customer-prices/{price_id}", response_model=CustomerPriceRead)
async def update_customer_price(
    price_id: UUID,
    body: CustomerPriceUpdate,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_permission(CATALOG_MANAGE)),
) -> CustomerPriceRead:
    """Edita o valor de um preço especial (mesmo tenant obrigatório)."""
    repo = CustomerPriceRepository(db)
    cp = await repo.get(price_id)
    if not cp:
        raise NotFoundError("Preço especial não encontrado.")
    await repo.update(cp, {"price": body.price})
    await record_audit(
        db, action="update", entity="customer_price",
        entity_id=cp.id, user_id=user.id, tenant_id=user.tenant_id,
    )
    await db.commit()
    return cp

@router.delete("/customer-prices/{price_id}", status_code=204)
async def delete_customer_price(
    price_id: UUID,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_permission(CATALOG_MANAGE)),
) -> None:
    """Remove um preço especial (mesmo tenant obrigatório)."""
    repo = CustomerPriceRepository(db)
    cp = await repo.get(price_id)
    if not cp:
        raise NotFoundError("Preço especial não encontrado.")
    await repo.delete(cp)
    await record_audit(
        db, action="delete", entity="customer_price",
        entity_id=cp.id, user_id=user.id, tenant_id=user.tenant_id,
    )
    await db.commit()

# ---------- Importação em massa de preços especiais (Bloco B3) ----------
@router.post("/customer-prices/import", response_model=CustomerPriceImportResult)
async def import_customer_prices(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_permission(CATALOG_MANAGE)),
) -> CustomerPriceImportResult:
    """Importa preços especiais em massa (CSV ou Excel).

    Colunas esperadas:
      - document: CPF/CNPJ do cliente (identificação estável).
      - sku: SKU do produto.
      - price: preço especial (aceita vírgula decimal).
    Par cliente+produto já existente → atualiza o preço (substitui valor).
    """
    content = await file.read()
    if len(content) > IMPORT_MAX_BYTES:
        raise ValidationFailedError("Arquivo excede o limite de 2 MB.")

    product_repo = ProductRepository(db)
    customer_repo = CustomerRepository(db)
    price_repo = CustomerPriceRepository(db)

    created = 0
    updated = 0
    skipped = 0
    errors: list[dict] = []
    # Evita duplicar par (cliente, produto) repetido no MESMO arquivo
    resolved: dict[tuple[UUID, UUID], CustomerPrice | None] = {}

    try:
        for row in _iter_price_rows(file.filename or "", content):
            document = _normalize_digits(row.get("document") or "")
            sku = (row.get("sku") or "").strip()
            price_text = (row.get("price") or "").strip().replace(",", ".")
            try:
                price = Decimal(price_text)
            except Exception:  # noqa: BLE001
                price = Decimal("0")
            if price <= 0:
                skipped += 1
                errors.append({"row": row, "error": "Preço deve ser um número maior que zero."})
                continue

            customer = await customer_repo.get_by_document(document) if document else None
            if customer is None:
                skipped += 1
                errors.append({"row": row, "error": "Cliente não encontrado para o documento informado."})
                continue

            product = await product_repo.get_by_sku(sku) if sku else None
            if product is None:
                skipped += 1
                errors.append({"row": row, "error": "Produto não encontrado para o SKU informado."})
                continue

            key = (customer.id, product.id)
            if key in resolved:
                # Repetido no próprio arquivo → também atualiza o preço
                existing = resolved[key]
            else:
                existing = await price_repo.get_for_product(customer.id, product.id)
                resolved[key] = existing

            if existing is not None:
                await price_repo.update(existing, {"price": price})
                updated += 1
            else:
                new_cp = await price_repo.create(
                    {"customer_id": customer.id, "product_id": product.id, "price": price}
                )
                resolved[key] = new_cp
                created += 1

        await db.commit()
    except Exception as exc:  # noqa: BLE001
        await db.rollback()
        raise ValidationFailedError(f"Falha ao processar o arquivo: {exc}")

    await record_audit(
        db, action="import", entity="customer_price",
        user_id=user.id, tenant_id=user.tenant_id,
    )
    await db.commit()
    return CustomerPriceImportResult(
        created=created, updated=updated, skipped=skipped, errors=errors
    )

# ---------- Cotação de preço (seção 17) ----------
@router.get("/products/{product_id}/quote", response_model=PriceQuote)
async def quote_product_price(
    product_id: UUID,
    customer_id: UUID | None = None,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
) -> PriceQuote:
    """Calcula o preço final de um produto para um cliente (seção 17).

    O backend recalcula o preço — nunca confia no frontend.
    Prioridade: preço do cliente > preço padrão do produto.
    Cliente do portal: usa SEMPRE o próprio customer_id (anti-vazamento).
    """
    repo = ProductRepository(db)
    product = await repo.get(product_id)
    if not product:
        raise NotFoundError("Produto não encontrado.")
    effective_customer = _resolve_quote_customer(user, customer_id)
    final_price, source = await calculate_product_price(db, product, effective_customer)
    return PriceQuote(
        product_id=product.id,
        sku=product.sku,
        name=product.name,
        base_price=product.price,
        customer_price=final_price if source == "customer" else None,
        final_price=final_price,
        price_source=source,
    )